"""Importa i programmi teorici (MTT) dagli Excel ufficiali.

    py -3.11 scripts/import_programma_mtt.py ["docs/sorgenti/Programmi Type Training/CH-47F - MTT - B1.3.xlsx" ...]

Senza argomenti importa tutti i file "<MDS> - MTT - <categoria>.xlsx" della cartella
`docs/sorgenti/Programmi Type Training`. Per ciascuno genera `src/dati/programmi/<id>.json`:
moduli e **materie**, cioè i gruppi di voci che condividono gli stessi "Tuition Min." del
programma (nell'Excel il valore compare sulla prima riga del gruppo e le righe successive lo
ereditano). Ogni materia è l'unità di lezione: ha una durata in minuti, un titolo e i chapter che copre.
"""

import json
import re
import sys
from pathlib import Path

import openpyxl

RADICE = Path(__file__).resolve().parent.parent
CARTELLA = RADICE / "docs/sorgenti/Programmi Type Training"
NOME_FILE = re.compile(r"^(?P<mds>.+?) - (?P<tipo>MTT|PTR) - (?P<categoria>.+)$")


def pulisci(v):
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def codice_chapter(v):
    return f"{v:02d}" if isinstance(v, int) else pulisci(v)


def minuti(v):
    """Nel file i totali di modulo oltre le mille unità sono scritti come 3.12 invece di 3120."""
    if isinstance(v, float) and not v.is_integer():
        return round(v * 1000)
    return int(v)


def descrivi(percorso):
    """MDS, tipo, categoria e id del programma dal nome del file ("CH-47F - MTT - B1.3.xlsx")."""
    m = NOME_FILE.match(Path(percorso).stem)
    assert m, f"Nome file inatteso (serve '<MDS> - MTT|PTR - <categoria>.xlsx'): {percorso}"
    mds, tipo, categoria = m["mds"].strip(), m["tipo"], m["categoria"].strip()
    sigla = lambda s: re.sub(r"[^a-z0-9]", "", s.lower())
    return {"mds": mds, "tipo": tipo, "categoria": categoria, "id": f"{tipo.lower()}-{sigla(mds)}-{sigla(categoria)}"}


def leggi_mtt(percorso):
    """Moduli, materie e tutte le voci (con il modulo) del foglio MTT."""
    libro = openpyxl.load_workbook(percorso, read_only=True)
    foglio = libro["MTT"] if "MTT" in libro.sheetnames else libro.active
    righe = list(foglio.iter_rows(values_only=True))
    intestazione = [pulisci(c).upper() for c in righe[0][:5]]
    assert intestazione in (["ITEM", "CH.", "SUBJECT", "LEVEL", "TUITION MIN."], ["ITEM", "ATA CH.", "SUBJECT", "LEVEL", "TUITION MIN."]), f"Intestazione inattesa: {intestazione}"

    moduli, materie, voci, attesa = [], [], [], []
    modulo, gruppo = None, None
    for r in righe[1:]:
        prima = pulisci(r[0])
        if prima.upper().startswith("TOTAL"):
            continue
        if prima.lower().startswith("module"):
            modulo = int(re.search(r"\d+", prima).group())
            moduli.append({"numero": modulo, "titolo": pulisci(r[2]).rstrip(":"), "minutiDichiarati": minuti(r[4]) if r[4] is not None else None})
            attesa, gruppo = [], None  # voci "coperte altrove" (livello -) rimaste senza durata nel modulo precedente
            continue
        if prima.upper().startswith("AVES"):
            # riga di gruppo dei moduli AVES ("AVES 2 – A.G.E. OF GENERAL USE"): la sua durata vale
            # per le voci che seguono solo se queste non hanno una durata propria
            gruppo = {"minuti": minuti(r[4]) if r[4] is not None else 0, "titolo": pulisci(r[2])}
            continue
        if r[0] is None or modulo is None:
            continue
        voce = {"item": int(r[0]), "chapter": codice_chapter(r[1]), "subject": pulisci(r[2]), "livello": pulisci(r[3])}
        voci.append({**voce, "modulo": modulo})
        if r[4] is None and gruppo and gruppo["minuti"]:
            materie.append({"id": f"m{len(materie) + 1:03d}", "modulo": modulo, "minuti": gruppo["minuti"], "titolo": gruppo["titolo"] or voce["subject"], "voci": [voce]})
            gruppo = None
            continue
        gruppo = None
        if r[4] is not None and minuti(r[4]) > 0:  # la durata apre una nuova materia, che assorbe le voci rimaste in attesa
            materie.append({"id": f"m{len(materie) + 1:03d}", "modulo": modulo, "minuti": minuti(r[4]), "titolo": (attesa[0] if attesa else voce)["subject"], "voci": [*attesa, voce]})
            attesa = []
        elif materie and materie[-1]["modulo"] == modulo:
            materie[-1]["voci"].append(voce)  # voce senza durata: fa parte della materia precedente
        else:
            attesa.append(voce)  # a inizio modulo: aspetta la prima durata

    for m in moduli:
        m["minuti"] = sum(x["minuti"] for x in materie if x["modulo"] == m["numero"])
        if m["minutiDichiarati"] not in (None, m["minuti"]):
            print(f"  ATTENZIONE modulo {m['numero']}: dichiarati {m['minutiDichiarati']} min, calcolati {m['minuti']} min")
        del m["minutiDichiarati"]
    for materia in materie:
        materia["chapters"] = [v["chapter"] for v in materia["voci"]]
    return moduli, materie, voci


def importa(percorso):
    d = descrivi(percorso)
    moduli, materie, _ = leggi_mtt(percorso)
    programma = {
        "id": d["id"],
        "tipo": "teorico",
        "nome": f"MTT {d['mds']} Cat. {d['categoria']}",
        "documento": f"T1 Military Type Training {d['mds']} Cat. {d['categoria']} – programma teorico (MTT)",
        "aeromobile": d["mds"],
        "categoria": d["categoria"],
        "moduli": moduli,
        "materie": materie,
    }
    (RADICE / "src/dati/programmi").mkdir(parents=True, exist_ok=True)
    (RADICE / f"src/dati/programmi/{d['id']}.json").write_text(json.dumps(programma, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    totale = sum(m["minuti"] for m in moduli)
    print(f"{d['id']}: {len(materie)} materie, {len(moduli)} moduli, {sum(len(m['voci']) for m in materie)} voci, {totale / 60:g} ore")


if __name__ == "__main__":
    for f in sys.argv[1:] or sorted(CARTELLA.glob("* - MTT - *.xlsx")):
        importa(f)
